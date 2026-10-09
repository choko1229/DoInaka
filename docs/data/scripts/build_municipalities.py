"""総務省「全国地方公共団体コード」Excel から municipalities_jp.csv を作る。

使い方: python3 -I build_municipalities.py <xlsx> <out_csv>
"""
from __future__ import annotations

import csv
import re
import sys
import unicodedata
from collections import Counter

import openpyxl

KIND_ORDER = ["pref", "designated_city", "city", "ward", "special_ward", "town", "village"]


def norm_code(v: object) -> str:
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    if not re.fullmatch(r"\d{5,6}", s):
        raise ValueError(f"bad code: {v!r}")
    return s.zfill(6)


def check_digit_ok(code: str) -> bool:
    # 上5桁に 6,5,4,3,2 を掛けた和を11で割った余り r。検査数字は (11 - r) の1の位
    d = [int(c) for c in code[:5]]
    total = sum(x * w for x, w in zip(d, [6, 5, 4, 3, 2]))
    return int(code[5]) == (11 - total % 11) % 10


def to_fullwidth_kana(s: str | None) -> str:
    if not s:
        return ""
    return unicodedata.normalize("NFKC", str(s)).strip()


def clean(s: object) -> str:
    if s is None:
        return ""
    return str(s).strip()


def main() -> None:
    xlsx, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    s1, s2 = wb.worksheets[0], wb.worksheets[1]
    print("sheet1:", s1.title, " sheet2:", s2.title)

    rows1 = [r for r in s1.iter_rows(values_only=True)][1:]
    rows2 = [r for r in s2.iter_rows(values_only=True)][1:]

    # 指定都市（シート2で 区 で終わらない行）
    designated: set[str] = set()
    wards: list[tuple] = []
    skipped2 = []
    for r in rows2:
        if r[0] is None:
            if any(x is not None for x in r):
                skipped2.append(r)
            continue
        code = norm_code(r[0])
        city = clean(r[2])
        if city.endswith("区"):
            wards.append(r)
        else:
            designated.add(code)

    records: list[dict[str, str]] = []
    skipped1 = []
    extra_cols = []
    for r in rows1:
        if r[0] is None:
            if any(x is not None for x in r):
                skipped1.append(r)
            continue
        if any(x not in (None, "") for x in r[5:]):
            extra_cols.append(r)
        code = norm_code(r[0])
        pref, city = clean(r[1]), clean(r[2])
        pk, ck = to_fullwidth_kana(r[3]), to_fullwidth_kana(r[4])
        if not city:
            kind = "pref"
        elif code in designated:
            kind = "designated_city"
        elif pref == "東京都" and city.endswith("区"):
            kind = "special_ward"
        elif city.endswith("市"):
            kind = "city"
        elif city.endswith("町"):
            kind = "town"
        elif city.endswith("村"):
            kind = "village"
        else:
            raise ValueError(f"unknown suffix: {r}")
        records.append(dict(code=code, pref_code=code[:2], pref_name=pref, city_name=city,
                            pref_kana=pk, city_kana=ck, kind=kind))

    for r in wards:
        code = norm_code(r[0])
        records.append(dict(code=code, pref_code=code[:2], pref_name=clean(r[1]), city_name=clean(r[2]),
                            pref_kana=to_fullwidth_kana(r[3]), city_kana=to_fullwidth_kana(r[4]), kind="ward"))

    records.sort(key=lambda d: d["code"])

    # ---- 検証 ----
    codes = [d["code"] for d in records]
    assert len(codes) == len(set(codes)), "duplicate code"
    cnt = Counter(d["kind"] for d in records)
    print("counts by kind:", {k: cnt.get(k, 0) for k in KIND_ORDER})
    shichoson = cnt["designated_city"] + cnt["city"] + cnt["town"] + cnt["village"]
    print("市町村 =", shichoson, " 特別区 =", cnt["special_ward"], " 市区町村 =", shichoson + cnt["special_ward"])
    print("市 (指定都市含む) =", cnt["designated_city"] + cnt["city"])
    print("designated cities in sheet2 =", len(designated), " missing in sheet1:",
          sorted(designated - {d['code'] for d in records if d['kind'] == 'designated_city'}))
    print("skipped non-code rows sheet1:", skipped1)
    print("skipped non-code rows sheet2:", skipped2)
    print("rows with extra columns:", extra_cols)
    bad_cd = [c for c in codes if not check_digit_ok(c)]
    print("check digit mismatches:", bad_cd)
    pref_rows = [d for d in records if d["kind"] == "pref"]
    print("pref rows:", len(pref_rows), " distinct pref_code:", len({d['pref_code'] for d in records}))

    # 末尾かなと種別の照合
    kana_suffix = {"city": ("シ",), "designated_city": ("シ",), "town": ("チョウ", "マチ"),
                   "village": ("ソン", "ムラ"), "ward": ("ク",), "special_ward": ("ク",)}
    mism = [d for d in records if d["kind"] in kana_suffix and not d["city_kana"].endswith(kana_suffix[d["kind"]])]
    print("kana suffix mismatches:", [(d["code"], d["city_name"], d["city_kana"]) for d in mism])
    town_m = Counter(d["city_kana"].endswith("マチ") for d in records if d["kind"] == "town")
    vil_m = Counter(d["city_kana"].endswith("ムラ") for d in records if d["kind"] == "village")
    print("town ending マチ:", town_m[True], " チョウ:", town_m[False])
    print("village ending ムラ:", vil_m[True], " ソン:", vil_m[False])
    # 名前の途中に別種別の字を含むもの（参考）
    inner = [d for d in records if d["kind"] in ("city", "town", "village", "designated_city")
             and re.search(r"[市町村区]", d["city_name"][:-1])]
    print("names containing 市/町/村/区 before suffix:", len(inner))
    for d in inner:
        print("  INNER", d["code"], d["city_name"], d["city_kana"], d["kind"])
    nonkana = [d for d in records if re.search(r"[^゠-ヿ]", d["city_kana"] + d["pref_kana"])]
    print("kana with non-katakana chars:", [(d["code"], d["city_kana"]) for d in nonkana])
    kagawa = [d for d in records if d["pref_code"] == "37"]
    print("Kagawa rows:", len(kagawa))
    for d in kagawa:
        print("  ", d)

    with open(out, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["code", "pref_code", "pref_name", "city_name", "pref_kana",
                                          "city_kana", "kind"], lineterminator="\n")
        w.writeheader()
        w.writerows(records)
    print("written", len(records), "rows")


if __name__ == "__main__":
    main()
